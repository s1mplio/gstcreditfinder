import csv
import io
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from gst_credit_finder.web import app

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "tests" / "fixtures" / "data"

# Credit at risk in the generated test data (both GSTR-2B months).
EXPECTED_AT_RISK = 39741.0

client = TestClient(app)


def signed_in_client(login="9876500001"):
    """A client logged in as a shopkeeper; returns (client, business_id)."""
    session = TestClient(app)
    response = session.post(
        "/api/signup",
        json={
            "name": "Test Owner",
            "login": login,
            "password": "correct horse",
            "role": "shop",
            "business_name": "Test Store",
        },
    )
    assert response.status_code == 200, response.text
    return session, response.json()["businesses"][0]["id"]


def gstr2b_files():
    return [
        ("gstr2b", (path.name, path.read_bytes(), "application/json"))
        for path in sorted(SAMPLE.glob("gstr2b_*.json"))
    ]


def test_home_page_and_assets():
    assert "Find the GST credit you" in client.get("/").text
    assert client.get("/static/app.js").status_code == 200
    assert client.get("/static/style.css").status_code == 200


def test_upload_csv_register():
    session, business_id = signed_in_client("9876500001")
    response = session.post(
        "/api/check",
        files=[("register", ("bills.csv", (SAMPLE / "purchase_register.csv").read_bytes(), "text/csv"))]
        + gstr2b_files(),
        data={"business_id": business_id},
    )

    assert response.status_code == 200, response.text
    data = response.json()
    assert data["buyer_name"] == "Test Store"
    assert data["total_at_risk"] == EXPECTED_AT_RISK
    assert {c["status"] for c in data["counts"]} >= {"matched", "missing_in_2b", "invalid_gstin"}
    assert data["reminders"][0]["supplier_name"] == "Sharma Traders"
    assert data["urgent"]

    report = session.get(f"/api/checks/{data['check_id']}/report")
    assert report.status_code == 200
    assert load_workbook(io.BytesIO(report.content)).sheetnames[0] == "Summary"


def test_upload_excel_register_with_title_row_and_other_column_names():
    with open(SAMPLE / "purchase_register.csv", encoding="utf-8") as file:
        rows = list(csv.reader(file))

    renamed = {
        "supplier_gstin": "GSTIN of Supplier",
        "invoice_number": "Invoice No.",
        "invoice_date": "Invoice Date",
        "taxable_value": "Taxable Value",
        "igst": "Integrated Tax",
        "cgst": "Central Tax",
        "sgst": "State/UT Tax",
    }

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Purchase Register September 2026"])  # title row
    sheet.append([renamed.get(name, name) for name in rows[0]])

    for row in rows[1:]:
        values = list(row)
        values[4:] = [float(value) for value in values[4:]]
        sheet.append(values)

    buffer = io.BytesIO()
    workbook.save(buffer)

    session, business_id = signed_in_client("9876500002")
    response = session.post(
        "/api/check",
        files=[("register", ("register.xlsx", buffer.getvalue(), "application/octet-stream"))]
        + gstr2b_files(),
        data={"business_id": business_id},
    )

    assert response.status_code == 200, response.text
    assert response.json()["total_at_risk"] == EXPECTED_AT_RISK


def test_register_missing_columns_is_a_clear_error():
    session, business_id = signed_in_client("9876500003")
    response = session.post(
        "/api/check",
        files=[("register", ("bills.csv", b"name,amount\nx,1\n", "text/csv"))] + gstr2b_files(),
        data={"business_id": business_id},
    )
    assert response.status_code == 400
    assert "missing columns" in response.json()["detail"]


def test_wrong_gstr2b_file_is_a_clear_error():
    session, business_id = signed_in_client("9876500004")
    response = session.post(
        "/api/check",
        files=[
            ("register", ("bills.csv", (SAMPLE / "purchase_register.csv").read_bytes(), "text/csv")),
            ("gstr2b", ("gstr2b.pdf", b"%PDF-1.4 not json", "application/pdf")),
        ],
        data={"business_id": business_id},
    )
    assert response.status_code == 400
    assert "PDF version" in response.json()["detail"]


def test_no_public_demo_and_template():
    assert client.get("/api/demo").status_code == 404
    assert client.get("/api/report/anything").status_code == 404

    template = client.get("/api/template")
    assert template.status_code == 200
    assert template.text.startswith("supplier_gstin,")
