import json
import subprocess
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from gst_credit_finder.cli import main
from gst_credit_finder.gstr2b import load_gstr2b

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "sample_data"


def test_sample_demo_finds_every_situation(tmp_path, capsys):
    subprocess.run([sys.executable, str(SAMPLE / "generate_sample.py")], check=True)

    out = tmp_path / "report.xlsx"

    code = main([
        "--register", str(SAMPLE / "purchase_register.csv"),
        "--gstr2b", str(SAMPLE / "gstr2b_092026.json"), str(SAMPLE / "gstr2b_102026.json"),
        "--buyer-name", "Shree Ganesh Kirana Stores",
        "--today", "2026-10-23",
        "--out", str(out),
    ])

    assert code == 0
    captured = capsys.readouterr()

    for label in (
        "Missing in GSTR-2B", "Amount mismatch", "IGST vs CGST/SGST mismatch",
        "Credit not available", "Invalid GSTIN on bill", "Probable match",
        "Duplicate bill", "In GSTR-2B, no bill", "URGENT",
    ):
        assert label in captured.out, label

    assert "invoice value is" in captured.err      # misread bill total warned
    assert "cdnr record(s) not checked" in captured.err

    workbook = load_workbook(out)
    assert workbook.sheetnames == ["Summary", "Action needed", "Supplier reminders", "All invoices"]
    assert workbook["Summary"]["B5"].value > 0
    assert "Sharma Traders" in workbook["Supplier reminders"]["A2"].value


def test_gstr2b_item_level_amounts_and_amendments(tmp_path):
    path = tmp_path / "2b.json"
    path.write_text(json.dumps({
        "data": {
            "gstin": "27AAPFU0939F1ZV",
            "rtnprd": "092026",
            "docdata": {
                "b2b": [{
                    "ctin": "27AAACR5055K1Z7", "trdnm": "SUPPLIER",
                    "inv": [
                        {"inum": "A-1", "dt": "01-09-2026", "itms": [
                            {"itm_det": {"txval": 100, "camt": 9, "samt": 9}},
                            {"itm_det": {"txval": 50, "camt": 4.5, "samt": 4.5}},
                        ]},
                        {"inum": "A-2", "dt": "02-09-2026", "txval": 200, "igst": 36},
                    ],
                }],
                "b2ba": [{
                    "ctin": "27AAACR5055K1Z7", "trdnm": "SUPPLIER",
                    "inv": [{"oinum": "A-2", "inum": "A-2R", "dt": "02-09-2026", "txval": 210, "igst": 37.8}],
                }],
            },
        },
    }))

    loaded = load_gstr2b(path)
    invoices = {inv.invoice_number: inv for inv in loaded.invoices}

    assert set(invoices) == {"A-1", "A-2R"}
    assert invoices["A-1"].taxable_value == Decimal("150.00")
    assert invoices["A-1"].total_tax == Decimal("27.00")
    assert invoices["A-1"].invoice_date == date(2026, 9, 1)
    assert invoices["A-2R"].igst == Decimal("37.80")
    assert loaded.period == "092026"
