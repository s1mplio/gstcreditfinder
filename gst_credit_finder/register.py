"""Load the shop's purchase register (bills it received) from CSV or Excel.

Required columns: supplier_gstin, invoice_number, invoice_date,
taxable_value. Optional: supplier_name, igst, cgst, sgst, cess,
invoice_value. Dates may be DD-MM-YYYY, DD/MM/YYYY, YYYY-MM-DD or
real Excel dates. Column names are matched case-insensitively and
common variants ("GSTIN of supplier", "Invoice No.") are recognised.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Dict, Iterable, List

from .gstin import clean as clean_gstin
from .models import Invoice, money

REQUIRED = ("supplier_gstin", "invoice_number", "invoice_date", "taxable_value")

DATE_FORMATS = ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y", "%d-%b-%Y", "%d %b %Y")

# Rounding allowed when checking taxable value + tax = invoice value.
TOTAL_TOLERANCE = Decimal("1.00")

# Common column names in accounting exports -> our column name.
ALIASES = {
    "gstin": "supplier_gstin",
    "supplier gstin": "supplier_gstin",
    "gstin of supplier": "supplier_gstin",
    "party gstin": "supplier_gstin",
    "supplier": "supplier_name",
    "supplier name": "supplier_name",
    "party name": "supplier_name",
    "trade name": "supplier_name",
    "invoice no": "invoice_number",
    "invoice number": "invoice_number",
    "bill no": "invoice_number",
    "invoice date": "invoice_date",
    "bill date": "invoice_date",
    "date": "invoice_date",
    "taxable value": "taxable_value",
    "taxable amount": "taxable_value",
    "integrated tax": "igst",
    "central tax": "cgst",
    "state tax": "sgst",
    "state/ut tax": "sgst",
    "utgst": "sgst",
    "invoice value": "invoice_value",
    "total": "invoice_value",
}


def parse_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = str(value or "").strip()

    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    raise ValueError(f"unrecognised date {text!r}")


def _column(name) -> str:
    key = re.sub(r"[^a-z0-9/ ]+", " ", str(name or "").lower())
    key = re.sub(r"\s+", " ", key).strip()

    if key.replace(" ", "_") in REQUIRED + ("supplier_name", "igst", "cgst", "sgst", "cess", "invoice_value"):
        return key.replace(" ", "_")

    return ALIASES.get(key, key.replace(" ", "_"))


@dataclass
class RegisterLoad:
    invoices: List[Invoice] = field(default_factory=list)
    # Rows that could not be read at all.
    errors: List[str] = field(default_factory=list)
    # Rows read, but whose amounts don't add up (likely misread).
    warnings: List[str] = field(default_factory=list)


def _text(value) -> str:
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        # Excel stores invoice "458" as 458.0
        return str(int(value))

    return str(value).strip()


def parse_rows(rows: Iterable[Dict], first_line: int = 2) -> RegisterLoad:
    """Build invoices from rows keyed by (any variant of) column name."""
    result = RegisterLoad()
    rows = list(rows)

    columns = {_column(name) for row in rows[:1] for name in row}

    missing = [name for name in REQUIRED if name not in columns]

    if missing:
        raise ValueError(
            "Purchase register is missing columns: "
            + ", ".join(missing)
            + ". Download the template to see the expected format."
        )

    for line, raw in enumerate(rows, start=first_line):
        row = {_column(key): value for key, value in raw.items() if key is not None}

        if not any(_text(value) for value in row.values()):
            continue

        try:
            invoice = Invoice(
                supplier_gstin=clean_gstin(_text(row["supplier_gstin"])),
                supplier_name=_text(row.get("supplier_name")),
                invoice_number=_text(row["invoice_number"]),
                invoice_date=parse_date(row["invoice_date"]),
                taxable_value=money(row["taxable_value"]),
                igst=money(row.get("igst")),
                cgst=money(row.get("cgst")),
                sgst=money(row.get("sgst")),
                cess=money(row.get("cess")),
                source="register",
                origin=f"row {line}",
            )
        except (ValueError, InvalidOperation) as error:
            result.errors.append(f"Row {line}: {error}")
            continue

        if not invoice.invoice_number:
            result.errors.append(f"Row {line}: invoice number is empty")
            continue

        invoice_value = _text(row.get("invoice_value"))

        if invoice_value:
            expected = invoice.taxable_value + invoice.total_tax

            if abs(money(invoice_value) - expected) > TOTAL_TOLERANCE:
                result.warnings.append(
                    f"Row {line} ({invoice.invoice_number}): taxable value + tax "
                    f"= {expected} but invoice value is {money(invoice_value)}; "
                    "check the bill"
                )

        result.invoices.append(invoice)

    return result


def _excel_rows(data: bytes) -> List[Dict]:
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)

    header = None

    # Skip title rows above the real header.
    for values in rows:
        names = {_column(value) for value in values if value is not None}

        if "supplier_gstin" in names and "invoice_number" in names:
            header = [str(value) if value is not None else None for value in values]
            break

    if header is None:
        raise ValueError(
            "Couldn't find the header row (needs supplier GSTIN and invoice number columns)."
        )

    return [dict(zip(header, values)) for values in rows]


def load_register_bytes(data: bytes, filename: str) -> RegisterLoad:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        return parse_rows(_excel_rows(data))

    text = data.decode("utf-8-sig", errors="replace")
    return parse_rows(csv.DictReader(io.StringIO(text)))


def load_register(path) -> RegisterLoad:
    return load_register_bytes(Path(path).read_bytes(), str(path))
