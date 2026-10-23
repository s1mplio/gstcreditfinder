"""Turn match results into a summary, supplier reminders and an Excel report."""

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .matcher import STATUSES, action, label, needs_action
from .models import ZERO, MatchResult

# Missing credit within this many days of its deadline is urgent.
URGENT_DAYS = 60

# Statuses where the supplier has to act.
SUPPLIER_ACTION = ("missing_in_2b", "amount_mismatch", "tax_type_mismatch")


def rupees(amount: Decimal) -> str:
    """Indian digit grouping: 1234567.5 -> '₹12,34,567.50'."""
    sign = "-" if amount < 0 else ""
    whole, _, fraction = f"{abs(amount):.2f}".partition(".")

    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        groups = []

        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]

        if head:
            groups.insert(0, head)

        whole = ",".join(groups + [tail])

    return f"{sign}₹{whole}.{fraction}"


@dataclass
class Summary:
    counts: Dict[str, int] = field(default_factory=dict)
    at_risk: Dict[str, Decimal] = field(default_factory=dict)
    total_at_risk: Decimal = ZERO
    claimable: Decimal = ZERO
    urgent: List[MatchResult] = field(default_factory=list)


def summarize(results: List[MatchResult], today: date) -> Summary:
    summary = Summary()

    for result in results:
        summary.counts[result.status] = summary.counts.get(result.status, 0) + 1
        summary.at_risk[result.status] = summary.at_risk.get(result.status, ZERO) + result.at_risk
        summary.total_at_risk += result.at_risk

        if result.status == "matched":
            summary.claimable += result.register.total_tax

        if result.deadline and (result.deadline - today).days <= URGENT_DAYS:
            summary.urgent.append(result)

    return summary


def supplier_reminders(results: List[MatchResult], buyer_name: str, buyer_gstin: str) -> Dict[str, dict]:
    """One ready-to-send message per supplier with something to fix."""
    by_supplier = defaultdict(list)

    for result in results:
        if result.status in SUPPLIER_ACTION:
            by_supplier[result.register.supplier_gstin].append(result)

    reminders = {}

    for supplier_gstin, items in by_supplier.items():
        name = items[0].register.supplier_name or supplier_gstin
        lines = []

        for result in items:
            bill = result.register
            problem = {
                "missing_in_2b": "not showing in our GSTR-2B",
                "amount_mismatch": "reported with different amounts",
                "tax_type_mismatch": "reported with the wrong tax type (IGST vs CGST/SGST)",
            }[result.status]

            lines.append(
                f"- Invoice {bill.invoice_number} dated {bill.invoice_date:%d-%m-%Y}, "
                f"GST {rupees(bill.total_tax)}: {problem}"
            )

        total = sum((r.at_risk for r in items), ZERO)

        message = (
            f"Dear {name},\n\n"
            f"We are {buyer_name} (GSTIN {buyer_gstin}). The following invoices "
            "from you need attention in your GST filing (GSTR-1):\n\n"
            + "\n".join(lines)
            + f"\n\nTotal GST credit affected: {rupees(total)}.\n"
            "Please file or amend these at the earliest so we can claim the credit.\n\n"
            "Thank you."
        )

        reminders[supplier_gstin] = {
            "supplier_name": name,
            "invoices": len(items),
            "at_risk": total,
            "message": message,
        }

    return dict(sorted(reminders.items(), key=lambda item: -item[1]["at_risk"]))


def text_summary(summary: Summary, period_text: str, today: date) -> str:
    lines = [
        f"GST CREDIT CHECK  {period_text}",
        "=" * 52,
        f"Credit at risk:       {rupees(summary.total_at_risk)}",
        f"Credit confirmed:     {rupees(summary.claimable)}",
        "",
    ]

    for status in STATUSES:
        count = summary.counts.get(status, 0)

        if not count:
            continue

        risk = summary.at_risk.get(status, ZERO)
        risk_text = f"  ({rupees(risk)} at risk)" if risk else ""
        lines.append(f"  {label(status):<28}{count:>4}{risk_text}")

    if summary.urgent:
        lines.append("")
        lines.append(f"URGENT: {len(summary.urgent)} item(s) within {URGENT_DAYS} days of the claim deadline")

        for result in summary.urgent:
            bill = result.register
            lines.append(
                f"  {bill.supplier_name or bill.supplier_gstin}: {bill.invoice_number} "
                f"({rupees(result.at_risk)}) deadline {result.deadline:%d-%m-%Y}, "
                f"{(result.deadline - today).days} days left"
            )

    return "\n".join(lines)


# =========================================================
# EXCEL
# =========================================================

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
STATUS_FILLS = {
    "matched": "E2EFDA",
    "probable_match": "FFF2CC",
    "duplicate_bill": "FFF2CC",
    "only_in_2b": "DDEBF7",
}
RISK_FILL = "F8CBAD"


def _write_table(sheet, headers, rows, widths=None):
    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in rows:
        sheet.append(row)

    sheet.freeze_panes = "A2"

    for index, header in enumerate(headers, start=1):
        width = (widths or {}).get(header, max(12, min(45, len(header) + 4)))
        sheet.column_dimensions[get_column_letter(index)].width = width


def _result_row(result: MatchResult) -> list:
    bill, record = result.register, result.gstr2b
    invoice = result.invoice

    return [
        label(result.status),
        invoice.supplier_name,
        invoice.supplier_gstin,
        bill.invoice_number if bill else "",
        record.invoice_number if record else "",
        invoice.invoice_date,
        float(bill.taxable_value) if bill else None,
        float(bill.total_tax) if bill else None,
        float(record.taxable_value) if record else None,
        float(record.total_tax) if record else None,
        record.period if record else "",
        float(result.at_risk) if result.at_risk else None,
        result.deadline,
        action(result.status),
        "; ".join(result.notes),
        bill.origin if bill else (record.origin if record else ""),
    ]


RESULT_HEADERS = [
    "Status", "Supplier", "Supplier GSTIN", "Invoice no. (bill)",
    "Invoice no. (GSTR-2B)", "Invoice date", "Taxable (bill)", "GST (bill)",
    "Taxable (GSTR-2B)", "GST (GSTR-2B)", "GSTR-2B period", "Credit at risk",
    "Claim deadline", "What to do", "Details", "Source",
]

RESULT_WIDTHS = {
    "Status": 24, "Supplier": 26, "Supplier GSTIN": 18, "What to do": 50,
    "Details": 60, "Invoice no. (bill)": 18, "Invoice no. (GSTR-2B)": 18,
}


def _format_results_sheet(sheet, results: List[MatchResult]):
    money_columns = ("G", "H", "I", "J", "L")
    date_columns = ("F", "M")

    for row_index, result in enumerate(results, start=2):
        fill = RISK_FILL if result.at_risk > 0 else STATUS_FILLS.get(result.status)

        if fill:
            sheet[f"A{row_index}"].fill = PatternFill("solid", fgColor=fill)

        for column in money_columns:
            sheet[f"{column}{row_index}"].number_format = "#,##0.00"

        for column in date_columns:
            sheet[f"{column}{row_index}"].number_format = "DD-MM-YYYY"

        for column in ("N", "O"):
            sheet[f"{column}{row_index}"].alignment = Alignment(wrap_text=True, vertical="top")


def write_excel(
    path,
    results: List[MatchResult],
    summary: Summary,
    reminders: Dict[str, dict],
    buyer_name: str,
    buyer_gstin: str,
    period_text: str,
    today: date,
):
    workbook = Workbook()

    # ---------------- Summary ----------------
    sheet = workbook.active
    sheet.title = "Summary"

    sheet["A1"] = "GST Credit Check"
    sheet["A1"].font = Font(bold=True, size=16)
    sheet["A2"] = f"{buyer_name}  |  GSTIN {buyer_gstin}  |  {period_text}"
    sheet["A3"] = f"Prepared {today:%d-%m-%Y}"

    sheet["A5"] = "Credit at risk"
    sheet["B5"] = float(summary.total_at_risk)
    sheet["A6"] = "Credit confirmed in GSTR-2B"
    sheet["B6"] = float(summary.claimable)

    for cell in ("A5", "A6"):
        sheet[cell].font = Font(bold=True)

    sheet["B5"].font = Font(bold=True, color="C00000", size=13)

    for cell in ("B5", "B6"):
        sheet[cell].number_format = "₹#,##0.00"

    row = 8
    for header, column in (("Result", "A"), ("Invoices", "B"), ("Credit at risk", "C"), ("What to do", "D")):
        cell = sheet[f"{column}{row}"]
        cell.value = header
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    for status in STATUSES:
        count = summary.counts.get(status, 0)

        if not count:
            continue

        row += 1
        sheet[f"A{row}"] = label(status)
        sheet[f"B{row}"] = count
        sheet[f"C{row}"] = float(summary.at_risk.get(status, ZERO))
        sheet[f"C{row}"].number_format = "₹#,##0.00"
        sheet[f"D{row}"] = action(status)

    for column, width in (("A", 30), ("B", 14), ("C", 18), ("D", 80)):
        sheet.column_dimensions[column].width = width

    # ---------------- Action needed ----------------
    action_results = [r for r in results if needs_action(r.status)]
    sheet = workbook.create_sheet("Action needed")
    _write_table(sheet, RESULT_HEADERS, [_result_row(r) for r in action_results], RESULT_WIDTHS)
    _format_results_sheet(sheet, action_results)

    # ---------------- Supplier reminders ----------------
    sheet = workbook.create_sheet("Supplier reminders")
    _write_table(
        sheet,
        ["Supplier", "Supplier GSTIN", "Invoices", "Credit at risk", "Message to send"],
        [
            [info["supplier_name"], gstin, info["invoices"], float(info["at_risk"]), info["message"]]
            for gstin, info in reminders.items()
        ],
        {"Supplier": 26, "Supplier GSTIN": 18, "Message to send": 90},
    )

    for row_index in range(2, len(reminders) + 2):
        sheet[f"D{row_index}"].number_format = "₹#,##0.00"
        sheet[f"E{row_index}"].alignment = Alignment(wrap_text=True, vertical="top")

    # ---------------- All invoices ----------------
    sheet = workbook.create_sheet("All invoices")
    _write_table(sheet, RESULT_HEADERS, [_result_row(r) for r in results], RESULT_WIDTHS)
    _format_results_sheet(sheet, results)

    workbook.save(path)
